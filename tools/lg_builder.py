"""Builds an LSGH Learning Guide (.docx) from a content module.

Usage:
    python3 tools/lg_builder.py content/grade7_term1_lg.py output/grade7/<name>.docx

The content module defines HEADER, OVERVIEW, OUTCOMES, POLICIES, ASSESSMENT,
STANDARDS, UNITS, MARK, STAGES and SIGNATORIES (see content/grade7_term1_lg.py).

Lesson bodies are lists of blocks:
    ("cycle", text)            yellow-highlighted cycle/meeting range
    ("lesson", text)           lesson title line
    ("meet", text)             meeting sub-heading inside a lesson
    ("h", text)                bold heading, e.g. "INTRODUCTION:"
    ("p", text)                plain paragraph
    ("q", text)                quoted teacher script (italic)
    ("ts", text)               TRANSITION STATEMENT
    ("step", text, [subs])     numbered activity (numbering runs per lesson)
    ("bul", [items])           bullet list
    ("obj", [items], [pillar keys])  objectives, then the pillars they focus on
    ("ex", title, [lines])     shaded worked-example / expected-output box
    ("sim", {...})             SIM LAB / LIVE SIM card: kind, title, tool, link,
                               setup, mission [..], questions [..], discovery
    ("table", [header], [[row], ...])  small grid, e.g. a pattern table
    ("pillar", [keys], [reasons])    pillar focus with optional reasons
    ("res", [items])           RESOURCES list
    ("sep",)                   separator between lessons
Markup inside text: **bold**, ^^green bold^^, __italic__; "(ID#n)" tags are
rendered in red bold automatically.
"""

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.shared import Pt, Twips

sys.path.insert(0, str(Path(__file__).resolve().parent))
import cm_builder as cm  # noqa: E402

cm.FONT = "Avenir"
cm.GREEN = "538135"
TEMPLATE = cm.ROOT / "templates" / "lsgh_template_portrait.docx"
PAGE_WIDTH = 10512  # 8.5in minus 0.6in + 0.5in margins, less a small allowance
cm.PAGE_WIDTH = PAGE_WIDTH

GRAY = "D0CECE"
EX_FILL = "EEF5E9"
RED = "FF0000"
STD_COLS = [2300, 8212]
POL_COLS = [2300, 8212]
ASSESS_COLS = [2500, 1700, 6312]
PLAN_COLS = [1250, 9262]
ID_TAG = re.compile(r"(\(ID#[^)]*\))")


def runs(par, text, size=10, color=None, bold=False, italic=False):
    for piece in ID_TAG.split(text):
        if not piece:
            continue
        if ID_TAG.fullmatch(piece):
            cm.add_runs(par, piece, size=size, color=RED, bold=True)
        else:
            cm.add_runs(par, piece, size=size, color=color, bold=bold, italic=italic)


def par_in(container, first=False, space_after=3, indent=0, align=None, keep=False):
    par = cm.new_par(container, first)
    par.paragraph_format.space_after = Pt(space_after)
    if indent:
        par.paragraph_format.left_indent = Twips(indent)
    if align is not None:
        par.alignment = align
    if keep:
        par.paragraph_format.keep_with_next = True
    return par


def marker_par(container, marker, text, size=10, indent=0, first=False, italic=False,
               bold=False, space_after=2):
    par = par_in(container, first, space_after)
    hang = 260 if len(marker) <= 2 else 360
    par.paragraph_format.left_indent = Twips(indent + hang)
    par.paragraph_format.first_line_indent = Twips(-hang)
    par.paragraph_format.tab_stops.add_tab_stop(Twips(indent + hang))
    cm.add_runs(par, marker + "\t", size=size, bold=bold)
    runs(par, text, size=size, italic=italic, bold=bold)
    return par


def heading(doc, text, roman):
    par = doc.add_paragraph()
    par.paragraph_format.space_before = Pt(10)
    par.paragraph_format.space_after = Pt(4)
    par.paragraph_format.keep_with_next = True
    cm.add_runs(par, f"{roman}.\t{text}", size=10, bold=True)
    par.paragraph_format.tab_stops.add_tab_stop(Twips(500))


def body_par(doc, text, indent=500):
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    par.paragraph_format.left_indent = Twips(indent)
    par.paragraph_format.space_after = Pt(6)
    runs(par, text)


def new_table(doc, widths, border="000000"):
    table = doc.add_table(rows=0, cols=len(widths))
    cm.set_table_borders(table, color=border)
    cm.fix_table_layout(table, widths)
    return table


def add_row(table, widths):
    row = table.add_row()
    cm.set_widths(row, widths)
    for cell in row.cells:
        cm.set_cell_margins(cell, 70, 70, 110, 110)
    return row


def label_cell(cell, text, fill=None, color=None, size=10):
    if fill:
        cm.set_cell_fill(cell, fill)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for i, line in enumerate(text.split("\n")):
        par = par_in(cell, first=i == 0, space_after=0, align=WD_ALIGN_PARAGRAPH.CENTER)
        cm.add_runs(par, line, size=size, bold=i == 0 or fill is None, color=color)


# ---------------------------------------------------------------- front matter

def build_header(doc, header):
    for line in header["lines"]:
        par = doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.space_after = Pt(0)
        cm.add_runs(par, line, size=10, bold=True)
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.space_before = Pt(10)
    cm.add_runs(par, header["title"], size=10, bold=True)


def build_overview(doc, overview, outcomes):
    heading(doc, "COURSE OVERVIEW:", "I")
    for text in overview:
        body_par(doc, text)
    heading(doc, "LEARNING OUTCOME:", "II")
    for code, text in outcomes:
        par = doc.add_paragraph()
        par.paragraph_format.left_indent = Twips(860)
        par.paragraph_format.first_line_indent = Twips(-300)
        par.paragraph_format.space_after = Pt(4)
        cm.add_runs(par, "•  ", size=10)
        cm.add_runs(par, code + " ", size=10, bold=True, italic=True)
        cm.add_runs(par, text, size=10, italic=True)


def build_policies(doc, policies):
    heading(doc, "POLICIES AND PROCEDURES:", "III")
    table = new_table(doc, POL_COLS)
    for pol in policies:
        row = add_row(table, POL_COLS)
        label_cell(row.cells[0], pol["label"])
        cell = row.cells[1]
        first = True
        for n, item in enumerate(pol["items"], 1):
            text, subs = (item, []) if isinstance(item, str) else item
            if len(pol["items"]) == 1:
                runs(par_in(cell, first, 4, align=WD_ALIGN_PARAGRAPH.JUSTIFY), text)
            else:
                marker_par(cell, f"{n}.", text, first=first, space_after=4)
            for sub in subs:
                marker_par(cell, "•", sub, indent=360)
            first = False
        if pol.get("note"):
            runs(par_in(cell, False, 2), pol["note"], italic=True)


def build_assessment(doc, assess):
    heading(doc, "METHODS OF ASSESSMENT:", "IV")
    table = new_table(doc, ASSESS_COLS)
    head = add_row(table, ASSESS_COLS)
    cm.repeat_header(head)
    for cell, text in zip(head.cells, ["GRADE COMPONENTS", "PERCENTAGE",
                                       "ASSESSMENT/REQUIREMENTS"]):
        label_cell(cell, text, fill=cm.GREEN, color=cm.WHITE, size=9.5)
    for comp in assess:
        row = add_row(table, ASSESS_COLS)
        cm.no_split(row)
        label_cell(row.cells[0], comp["component"])
        label_cell(row.cells[1], comp["weight"])
        cell = row.cells[2]
        first = True
        for group, items in comp["groups"]:
            runs(par_in(cell, first, 1), f"**{group}**")
            first = False
            for item in items:
                runs(par_in(cell, False, 1, indent=300), item, size=9.5)


def build_standards(doc, standards):
    heading(doc, "LEARNING PLAN", "V")
    table = new_table(doc, STD_COLS)
    for block in standards:
        row = add_row(table, STD_COLS)
        label_cell(row.cells[0], block["label"], fill=cm.GREEN, color=cm.WHITE)
        cell = row.cells[1]
        first = True
        for kind, *data in block["body"]:
            if kind == "p":
                runs(par_in(cell, first, 4, align=WD_ALIGN_PARAGRAPH.JUSTIFY), data[0])
            elif kind == "lead":
                runs(par_in(cell, first, 2), data[0])
            elif kind == "h":
                runs(par_in(cell, first, 2, keep=True), f"**{data[0]}**", size=10.5)
            elif kind == "center":
                runs(par_in(cell, first, 4, align=WD_ALIGN_PARAGRAPH.CENTER),
                     f"**{data[0]}**", size=11)
            elif kind == "num":
                for n, item in enumerate(data[0], 1):
                    marker_par(cell, f"{n}.", item, first=first and n == 1)
            elif kind == "bul":
                for n, item in enumerate(data[0]):
                    text, subs = (item, []) if isinstance(item, str) else item
                    marker_par(cell, "•", text, first=first and n == 0)
                    for sub in subs:
                        marker_par(cell, "–", sub, indent=300, size=9.5)
            elif kind == "sub":
                runs(par_in(cell, first, 1, indent=200), f"**{data[0]}**")
                for item in data[1]:
                    runs(par_in(cell, False, 3, indent=400), item, size=9.5, italic=True)
            first = False


# ---------------------------------------------------------------- learning plan

def render_blocks(cell, blocks, first=True):
    step = 0
    for block in blocks:
        kind = block[0]
        if kind == "cycle":
            par = par_in(cell, first, 6, keep=True)
            par.paragraph_format.space_before = Pt(4)
            runs(par, block[1], bold=True)
            for run in par.runs:
                run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        elif kind == "lesson":
            step = 0
            runs(par_in(cell, first, 6, keep=True), block[1], bold=True)
        elif kind == "meet":
            par = par_in(cell, first, 3, keep=True)
            par.paragraph_format.space_before = Pt(6)
            runs(par, f"^^{block[1]}^^", size=10)
        elif kind == "h":
            par = par_in(cell, first, 2, keep=True)
            par.paragraph_format.space_before = Pt(6)
            runs(par, block[1], bold=True)
        elif kind == "p":
            runs(par_in(cell, first, 4, align=WD_ALIGN_PARAGRAPH.JUSTIFY), block[1])
        elif kind == "q":
            par = par_in(cell, first, 4, indent=200, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
            runs(par, f"“{block[1]}”", italic=True)
        elif kind == "ts":
            par = par_in(cell, first, 4, indent=360, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
            cm.add_runs(par, "TRANSITION STATEMENT: ", size=9.5, bold=True)
            runs(par, f"“{block[1]}”", size=9.5, italic=True)
        elif kind == "step":
            step += 1
            marker_par(cell, f"{step}.", block[1], first=first, space_after=3)
            for sub in (block[2] if len(block) > 2 else []):
                marker_par(cell, "•", sub, indent=360, size=9.5)
        elif kind == "bul":
            for n, item in enumerate(block[1]):
                text, subs = (item, []) if isinstance(item, str) else item
                marker_par(cell, "•", text, first=first and n == 0, indent=200)
                for sub in subs:
                    marker_par(cell, "–", sub, indent=560, size=9.5)
        elif kind == "obj":
            for n, item in enumerate(block[1]):
                marker_par(cell, "•", item, first=first and n == 0, indent=200)
            pillar_focus(cell, block[2], label="All objectives focus on the following:")
        elif kind == "pillar":
            pillar_focus(cell, block[1], reasons=block[2] if len(block) > 2 else None,
                         label="Pillar focus:")
        elif kind == "ex":
            example_box(cell, block[1], block[2])
        elif kind == "sim":
            sim_box(cell, block[1])
        elif kind == "table":
            grid(cell, block[1], block[2])
        elif kind == "res":
            par = par_in(cell, first, 2, keep=True)
            par.paragraph_format.space_before = Pt(6)
            cm.add_runs(par, "RESOURCES:", size=10, bold=True)
            for item in block[1]:
                marker_par(cell, "•", item, indent=0)
        elif kind == "sep":
            par = par_in(cell, first, 6, align=WD_ALIGN_PARAGRAPH.CENTER)
            par.paragraph_format.space_before = Pt(10)
            cm.add_runs(par, ">>>>>>>>>>>>>>>>>>>>=============<<<<<<<<<<<<<<<<<<<<",
                        size=9, bold=True)
        else:
            raise ValueError(f"unknown block {kind!r}")
        first = False


def pillar_focus(cell, keys, reasons=None, label=None):
    if label:
        runs(par_in(cell, False, 0, indent=360), label, size=8.5, italic=True)
    for key in keys:
        name, focus, indicator = cm.PILLARS[key]
        runs(par_in(cell, False, 0, indent=360), f"Pillar: {name} – {focus}", size=8.5)
        runs(par_in(cell, False, 3, indent=360), f"Indicator (G7–10): {indicator}", size=8.5)
    if reasons:
        runs(par_in(cell, False, 0, indent=360), "Reasons:", size=8.5, bold=True)
        for reason in reasons:
            marker_par(cell, "•", reason, indent=360, size=8.5, space_after=1)


def example_box(cell, title, lines):
    width = PLAN_COLS[1] - 700
    box = cell.add_table(rows=1, cols=1)
    cm.set_table_borders(box, color="A9C79A")
    cm.fix_table_layout(box, [width])
    cm.no_split(box.rows[0])
    inner = box.rows[0].cells[0]
    inner.width = Twips(width)
    cm.set_cell_fill(inner, EX_FILL)
    cm.set_cell_margins(inner, 80, 80, 140, 140)
    runs(par_in(inner, True, 3), f"^^{title}^^", size=9.5)
    for line in lines:
        if line.startswith("- "):
            marker_par(inner, "•", line[2:], size=9.5, space_after=1)
        elif line.startswith("  "):
            runs(par_in(inner, False, 1, indent=400), line.strip(), size=9.5)
        else:
            runs(par_in(inner, False, 2), line, size=9.5)
    cell.add_paragraph().paragraph_format.space_after = Pt(0)


def sim_box(cell, sim):
    """A simulation card: digital SIM LAB (blue) or LIVE SIM role-play (orange)."""
    live = sim.get("kind") == "LIVE SIM"
    fill, border, ink = (("FCE4D6", "F4B183", "843C0C") if live
                         else ("DDEBF7", "9DC3E6", "1F4E79"))
    width = PLAN_COLS[1] - 700
    box = cell.add_table(rows=1, cols=1)
    cm.set_table_borders(box, color=border, size=8)
    cm.fix_table_layout(box, [width])
    cm.no_split(box.rows[0])
    inner = box.rows[0].cells[0]
    inner.width = Twips(width)
    cm.set_cell_fill(inner, fill)
    cm.set_cell_margins(inner, 90, 90, 150, 150)
    par = par_in(inner, True, 3)
    cm.add_runs(par, f"{sim.get('kind', 'SIM LAB')}  ·  ", size=10, bold=True, color=ink)
    cm.add_runs(par, sim["title"], size=10, bold=True, color=ink)
    for label, key in (("Tool", "tool"), ("Link", "link"), ("Set-up", "setup")):
        if sim.get(key):
            runs(par_in(inner, False, 1), f"**{label}:** {sim[key]}", size=9)
    if sim.get("mission"):
        runs(par_in(inner, False, 1), "**Mission:**", size=9.5)
        for n, step in enumerate(sim["mission"], 1):
            marker_par(inner, f"{n}.", step, size=9.5, space_after=1)
    if sim.get("questions"):
        runs(par_in(inner, False, 1), "**Notice and Wonder:**", size=9.5)
        for q in sim["questions"]:
            marker_par(inner, "?", q, size=9.5, space_after=1)
    if sim.get("discovery"):
        runs(par_in(inner, False, 1), f"**Expected discovery:** __{sim['discovery']}__",
             size=9.5)
    cell.add_paragraph().paragraph_format.space_after = Pt(0)


def grid(cell, header, rows):
    width = PLAN_COLS[1] - 700
    widths = [width // len(header)] * len(header)
    table = cell.add_table(rows=0, cols=len(header))
    cm.set_table_borders(table, color="7F7F7F")
    cm.fix_table_layout(table, widths)
    for r, values in enumerate([header] + rows):
        row = table.add_row()
        cm.set_widths(row, widths)
        for c, (gcell, text) in enumerate(zip(row.cells, values)):
            cm.set_cell_margins(gcell, 30, 30, 60, 60)
            if r == 0:
                cm.set_cell_fill(gcell, cm.GREEN)
            par = par_in(gcell, True, 0, align=WD_ALIGN_PARAGRAPH.CENTER)
            runs(par, text, size=9, bold=r == 0, color=cm.WHITE if r == 0 else None)
    cell.add_paragraph().paragraph_format.space_after = Pt(0)


def build_plan(doc, units, mark, stages):
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    table = new_table(doc, PLAN_COLS)

    row = add_row(table, PLAN_COLS)
    cell = row.cells[0].merge(row.cells[1])
    runs(par_in(cell, True, 3), "**CHAPTER/UNIT, LESSON:**")
    for unit in units:
        runs(par_in(cell, False, 1, indent=100), f"**{unit['title']}**")
        for lesson in unit["lessons"]:
            title, subs = lesson
            runs(par_in(cell, False, 1, indent=300), f"**{title}**")
            for sub in subs:
                marker_par(cell, "•", sub, indent=500, size=9.5, space_after=0)
        cell.paragraphs[-1].paragraph_format.space_after = Pt(5)

    row = add_row(table, PLAN_COLS)
    cell = row.cells[0].merge(row.cells[1])
    runs(par_in(cell, True, 2), "**MARK: (LEARNING OBJECTIVES):** (ID#1)")
    runs(par_in(cell, False, 3),
         "At the end of the chapter/lesson, the students should be able to:")
    for code, text in mark:
        par = par_in(cell, False, 1)
        par.paragraph_format.left_indent = Twips(560)
        par.paragraph_format.first_line_indent = Twips(-560)
        par.paragraph_format.tab_stops.add_tab_stop(Twips(560))
        runs(par, f"^^{code}^^\t{text}", size=9.5)

    row = add_row(table, PLAN_COLS)
    cm.repeat_header(row)
    cell = row.cells[0].merge(row.cells[1])
    cm.set_cell_fill(cell, cm.GREEN)
    runs(par_in(cell, True, 0, align=WD_ALIGN_PARAGRAPH.CENTER), "LESSON DEVELOPMENT",
         color=cm.WHITE, bold=True)

    for stage in stages:
        row = add_row(table, PLAN_COLS)
        left, right = row.cells
        cm.set_cell_fill(left, GRAY)
        par = par_in(left, True, 2, align=WD_ALIGN_PARAGRAPH.CENTER)
        cm.add_runs(par, stage["name"], size=10, bold=True)
        par = par_in(left, False, 0, align=WD_ALIGN_PARAGRAPH.CENTER)
        cm.add_runs(par, stage["desc"], size=7.5)
        render_blocks(right, stage["blocks"])


def build_signatories(doc, sig):
    cm.build_signatories(doc, sig)


def build(content_path, out_path):
    c = cm.load_content(content_path)
    doc = Document(str(TEMPLATE))
    build_header(doc, c.HEADER)
    build_overview(doc, c.OVERVIEW, c.OUTCOMES)
    build_policies(doc, c.POLICIES)
    build_assessment(doc, c.ASSESSMENT)
    build_standards(doc, c.STANDARDS)
    build_plan(doc, c.UNITS, c.MARK, c.STAGES)
    build_signatories(doc, c.SIGNATORIES)
    out = Path(out_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    return out


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    print(build(sys.argv[1], sys.argv[2]))
