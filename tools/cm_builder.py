"""Builds an LSGH Curriculum Map (.docx) from a content module.

Usage:
    python3 tools/cm_builder.py courses/<course>/content/<module>.py courses/<course>/output/<name>.docx

The content module defines HEADER, STANDARDS, ROADMAP, ROWS and SIGNATORIES
(see courses/mathematics/content/grade7_term1_cm.py). Inline markup inside any text:
    **bold**          bold
    ^^text^^          bold, LSGH green
    __text__          italic
"""

import importlib.util
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Twips

ROOT = Path(__file__).resolve().parent.parent
TEMPLATE = ROOT / "templates" / "lsgh_template.docx"

FONT = "Arial"
GREEN = "38761D"
GREEN_DARK = "274E13"
WHITE = "FFFFFF"
FILL = {
    "label": GREEN,
    "banner": "D9EAD3",
    "quiz": "FFF2CC",
    "pt": "D0E0E3",
    "exam": "FCE5CD",
    "buffer": "EFEFEF",
    "roadmap_head": GREEN,
}

PAGE_WIDTH = 12960  # 11in landscape minus 1in margins, in twips
STD_COLS = [2600, 10360]
MAP_COLS = [1200, 1950, 2400, 1950, 2280, 1430, 1750]
MAP_HEADINGS = [
    "NO. OF MEETINGS",
    "UNIT TOPIC / CONTENT",
    "COMPETENCIES / SKILLS",
    "ASSESSMENTS",
    "ACTIVITIES",
    "RESOURCES",
    "PILLARS OF THE LSGH LASALLIAN IDENTITY",
]

PILLARS = {
    "SF": (
        "Spirit of Faith",
        "Striving to consistently live out our faith and values in our day-to-day choices.",
        "Uses critical thinking in evaluating moral dilemmas and applying ethical "
        "principles to real-life situations.",
    ),
    "ZS": (
        "Zeal for Service",
        "Striving to pursue excellence and integrity in our endeavors consistently.",
        "Demonstrates reflective learning and a growth mindset for global competence; "
        "practices accountability in decision-making.",
    ),
    "CM": (
        "Communion in Mission",
        "Developing open-mindedness, dialogue skills, and solidarity with all.",
        "Speaks confidently in sharing ideas and opinions on relevant issues while "
        "respecting differences; builds trust and confidence among peers.",
    ),
}

MARKUP = re.compile(r"(\*\*.+?\*\*|\^\^.+?\^\^|__.+?__)")


# ---------------------------------------------------------------- low level

def set_cell_fill(cell, hex_fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=60, bottom=60, left=80, right=80):
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for side, val in (("top", top), ("bottom", bottom), ("start", left), ("end", right)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tc_pr.append(mar)


def set_table_borders(table, color="000000", size=4):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tbl_pr.append(borders)


def fix_table_layout(table, widths):
    tbl_pr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for col, w in zip(grid.findall(qn("w:gridCol")), widths):
        col.set(qn("w:w"), str(w))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER


def set_widths(row, widths):
    for cell, w in zip(row.cells, widths):
        cell.width = Twips(w)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def no_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def add_runs(par, text, size=9, color=None, bold=False, italic=False):
    for piece in MARKUP.split(text):
        if not piece:
            continue
        b, i, c = bold, italic, color
        if piece.startswith("**") and piece.endswith("**"):
            piece, b = piece[2:-2], True
        elif piece.startswith("^^") and piece.endswith("^^"):
            piece, b, c = piece[2:-2], True, GREEN
        elif piece.startswith("__") and piece.endswith("__"):
            piece, i = piece[2:-2], True
        run = par.add_run(piece)
        run.font.name = FONT
        run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        run.font.size = Pt(size)
        run.bold = b
        run.italic = i
        if c:
            run.font.color.rgb = RGBColor.from_string(c)


def new_par(container, first=False):
    """Reuse the empty paragraph a fresh cell starts with."""
    if first and container.paragraphs and not container.paragraphs[0].text:
        par = container.paragraphs[0]
    else:
        par = container.add_paragraph()
    pf = par.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(2)
    pf.line_spacing = 1.0
    return par


def write_par(cell, text, first=False, size=9, align=None, color=None, bold=False,
              italic=False, marker=None, indent=0, space_after=2):
    par = new_par(cell, first)
    par.paragraph_format.space_after = Pt(space_after)
    if marker is not None:
        hang = 220 if len(marker) <= 2 else 330
        par.paragraph_format.left_indent = Twips(indent + hang)
        par.paragraph_format.first_line_indent = Twips(-hang)
        add_runs(par, marker + "\t", size=size, color=color, bold=bold)
        par.paragraph_format.tab_stops.add_tab_stop(Twips(indent + hang))
    elif indent:
        par.paragraph_format.left_indent = Twips(indent)
    add_runs(par, text, size=size, color=color, bold=bold, italic=italic)
    if align:
        par.alignment = align
    return par


def write_items(cell, items, first=False, numbered=False, size=9, bullet="•"):
    """items: str, or (str, [sub-items]) for a nested bullet list."""
    for n, item in enumerate(items, 1):
        text, subs = (item, []) if isinstance(item, str) else item
        marker = f"{n}." if numbered else bullet
        write_par(cell, text, first=first and n == 1, size=size, marker=marker)
        for sub in subs:
            write_par(cell, sub, size=size, marker="–", indent=220)


# ---------------------------------------------------------------- sections

def build_header(doc, header):
    for i, line in enumerate(header["lines"]):
        par = doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.space_after = Pt(0)
        add_runs(par, line, size=11, bold=True)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_runs(par, header["title"], size=12, bold=True)
    if header.get("note"):
        par = doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.space_after = Pt(6)
        add_runs(par, header["note"], size=8, italic=True, color="595959")


def build_standards(doc, standards):
    table = doc.add_table(rows=0, cols=2)
    set_table_borders(table)
    fix_table_layout(table, STD_COLS)
    for block in standards:
        row = table.add_row()
        set_widths(row, STD_COLS)
        label, body = row.cells
        set_cell_fill(label, FILL["label"])
        set_cell_margins(label, 100, 100, 100, 100)
        set_cell_margins(body, 100, 100, 140, 140)
        write_par(label, block["label"], first=True, size=10, bold=True, color=WHITE,
                  align=WD_ALIGN_PARAGRAPH.CENTER)
        first = True
        if block.get("lead"):
            write_par(body, block["lead"], first=True, size=10, space_after=4)
            first = False
        if block.get("items"):
            write_items(body, block["items"], first=first, numbered=block.get("numbered", False),
                        size=10)
            first = False
        for group in block.get("groups", []):
            write_par(body, group["title"], first=first, size=10, bold=True, space_after=2)
            first = False
            for gi in group["items"]:
                write_par(body, gi, size=9.5, italic=True, marker="–", indent=120)
        if block.get("text"):
            write_par(body, block["text"], first=first, size=10,
                      align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    return table


def build_roadmap(doc, roadmap):
    par = doc.add_paragraph()
    par.paragraph_format.page_break_before = True
    par.paragraph_format.space_after = Pt(4)
    add_runs(par, roadmap["title"], size=11, bold=True, color=GREEN_DARK)
    if roadmap.get("subtitle"):
        par = doc.add_paragraph()
        par.paragraph_format.space_after = Pt(4)
        add_runs(par, roadmap["subtitle"], size=9, italic=True)

    widths = roadmap["widths"]
    table = doc.add_table(rows=1, cols=len(widths))
    set_table_borders(table, color="7F7F7F")
    fix_table_layout(table, widths)
    head = table.rows[0]
    set_widths(head, widths)
    repeat_header(head)
    for cell, text in zip(head.cells, roadmap["columns"]):
        set_cell_fill(cell, FILL["roadmap_head"])
        set_cell_margins(cell)
        write_par(cell, text, first=True, size=9, bold=True, color=WHITE,
                  align=WD_ALIGN_PARAGRAPH.CENTER)
    for entry in roadmap["rows"]:
        row = table.add_row()
        set_widths(row, widths)
        no_split(row)
        fill = FILL.get(entry.get("kind", "banner"))
        for i, (cell, text) in enumerate(zip(row.cells, entry["cells"])):
            set_cell_margins(cell)
            if fill and i == 0:
                set_cell_fill(cell, fill)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            write_par(cell, text, first=True, size=9, bold=(i == 0))
    if roadmap.get("legend"):
        par = doc.add_paragraph()
        par.paragraph_format.space_before = Pt(4)
        add_runs(par, roadmap["legend"], size=8, italic=True, color="595959")


def pillar_cell(cell, keys):
    first = True
    for key in keys:
        name, focus, indicator = PILLARS[key]
        write_par(cell, f"^^Pillar: {name}^^ – {focus}", first=first, size=8.5)
        write_par(cell, f"**Indicator (G7–10):** {indicator}", size=8.5, space_after=6)
        first = False


def build_map(doc, rows):
    par = doc.add_paragraph()
    par.paragraph_format.page_break_before = True
    par.paragraph_format.space_after = Pt(0)
    table = doc.add_table(rows=1, cols=len(MAP_COLS))
    set_table_borders(table)
    fix_table_layout(table, MAP_COLS)
    head = table.rows[0]
    set_widths(head, MAP_COLS)
    repeat_header(head)
    for cell, text in zip(head.cells, MAP_HEADINGS):
        set_cell_fill(cell, GREEN)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        write_par(cell, text, first=True, size=9, bold=True, color=WHITE,
                  align=WD_ALIGN_PARAGRAPH.CENTER)

    for spec in rows:
        row = table.add_row()
        set_widths(row, MAP_COLS)
        if spec["type"] == "banner":
            no_split(row)
            cell = row.cells[0].merge(row.cells[-1])
            set_cell_fill(cell, FILL["banner"])
            set_cell_margins(cell, 80, 80, 140, 140)
            write_par(cell, spec["title"], first=True, size=11, bold=True, color=GREEN_DARK)
            write_par(cell, spec["sub"], size=9, color=GREEN_DARK)
            write_par(cell, f"**Big Question:** __{spec['question']}__", size=9.5)
            continue

        cells = row.cells
        for cell in cells:
            set_cell_margins(cell)
        fill = FILL.get(spec.get("kind"))
        if fill:
            for cell in cells:
                set_cell_fill(cell, fill)

        meet, topic, comps, assess, acts, res, pil = cells
        write_par(meet, spec["meetings"], first=True, size=9, bold=True)
        write_par(meet, spec["range"], size=9)

        write_par(topic, spec["topic"], first=True, size=9, bold=True)
        if spec.get("subtopics"):
            write_items(topic, spec["subtopics"])

        first = True
        if spec.get("comp_lead", True) and spec.get("comps"):
            write_par(comps, "The learners should be able to:", first=True, size=9)
            first = False
        for comp in spec.get("comps", []):
            if isinstance(comp, str):
                write_par(comps, comp, first=first, size=9, marker="•")
            else:
                code, deped, text = comp
                star = " ★" if deped else ""
                write_par(comps, f"^^{code}{star}^^  {text}", first=first, size=9)
            first = False

        first = True
        for kind, text in spec.get("assess", []):
            write_par(assess, f"**{kind}:** {text}", first=first, size=9, marker="•")
            first = False

        write_items(acts, spec.get("acts", []), first=True, numbered=True)
        write_items(res, spec.get("res", []), first=True)
        pillar_cell(pil, spec.get("pillars", []))
    return table


def build_signatories(doc, sig):
    par = doc.add_paragraph()
    par.paragraph_format.keep_with_next = True
    widths = [PAGE_WIDTH // len(sig)] * len(sig)
    table = doc.add_table(rows=1, cols=len(sig))
    fix_table_layout(table, widths)
    no_split(table.rows[0])
    set_widths(table.rows[0], widths)
    for cell, block in zip(table.rows[0].cells, sig):
        write_par(cell, block["label"], first=True, size=10, space_after=24)
        for name in block["names"]:
            write_par(cell, name, size=10, bold=True, space_after=0)
            write_par(cell, block["role"], size=10, space_after=18)


def load_content(path):
    spec = importlib.util.spec_from_file_location("cm_content", path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def build(content_path, out_path):
    c = load_content(content_path)
    doc = Document(str(TEMPLATE))
    build_header(doc, c.HEADER)
    build_standards(doc, c.STANDARDS)
    build_roadmap(doc, c.ROADMAP)
    build_map(doc, c.ROWS)
    build_signatories(doc, c.SIGNATORIES)
    out = Path(out_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    return out


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    print(build(sys.argv[1], sys.argv[2]))
